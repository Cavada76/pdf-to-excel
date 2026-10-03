#!/usr/bin/env python3
"""Extract tables from PDF files into a clean, formatted Excel workbook.

Usage:
    python pdf_to_excel.py input.pdf [more.pdf ...] -o output.xlsx
        [--pages 1-3,5] [--decimal auto|comma|dot] [--no-merge] [--csv]

Prints a JSON summary to stdout so the caller can report what was found.
Requires: pdfplumber, openpyxl  (pip install pdfplumber openpyxl)
"""
import argparse
import csv
import json
import re
import sys
from datetime import datetime
from pathlib import Path

try:
    import pdfplumber
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError as exc:  # pragma: no cover
    sys.exit(f"Missing dependency: {exc.name}. Run: pip install pdfplumber openpyxl")

TEXT_STRATEGY = {"vertical_strategy": "text", "horizontal_strategy": "text"}
CURRENCY_RE = re.compile(r"[$€£¥]|\b(?:USD|EUR|GBP|NOK|SEK|DKK|kr)\b", re.I)
DATE_FORMATS = ("%Y-%m-%d", "%d.%m.%Y", "%d.%m.%y")


def parse_pages(spec, total):
    if not spec:
        return list(range(total))
    pages = set()
    for part in spec.split(","):
        part = part.strip()
        if "-" in part:
            start, end = part.split("-", 1)
            pages.update(range(int(start), int(end) + 1))
        elif part:
            pages.add(int(part))
    return sorted(p - 1 for p in pages if 1 <= p <= total)


def parse_number(text, decimal="auto"):
    """Return a float for numeric-looking cells, else None."""
    s = CURRENCY_RE.sub("", text).replace(" ", " ").strip()
    negative = False
    if s.startswith("(") and s.endswith(")"):
        negative, s = True, s[1:-1].strip()
    if s.endswith("-") and len(s) > 1:  # trailing minus, common in bank statements
        negative, s = True, s[:-1].strip()
    if s.startswith(("-", "−")):
        negative, s = True, s[1:].strip()
    s = s.replace(" ", "").replace("'", "")
    if s.endswith("%"):
        s = s[:-1]
    if not s or not re.fullmatch(r"[\d.,]+", s) or not re.search(r"\d", s):
        return None

    if decimal == "comma":
        s = s.replace(".", "").replace(",", ".")
    elif decimal == "dot":
        s = s.replace(",", "")
    elif "," in s and "." in s:
        if s.rfind(",") > s.rfind("."):
            s = s.replace(".", "").replace(",", ".")
        else:
            s = s.replace(",", "")
    elif "," in s:
        s = s.replace(",", "") if re.fullmatch(r"\d{1,3}(,\d{3})+", s) else s.replace(",", ".")
    elif re.fullmatch(r"\d{1,3}(\.\d{3}){2,}", s):
        s = s.replace(".", "")
    try:
        value = float(s)
    except ValueError:
        return None
    return -value if negative else value


def parse_date(text):
    for fmt in DATE_FORMATS:
        try:
            return datetime.strptime(text.strip(), fmt).date()
        except ValueError:
            continue
    return None


def convert_cell(raw, decimal):
    if raw is None:
        return None
    text = " ".join(str(raw).split())
    if not text:
        return None
    if re.fullmatch(r"0\d+", text):  # keep leading zeros (account numbers, IDs)
        return text
    date = parse_date(text)
    if date:
        return date
    number = parse_number(text, decimal)
    return number if number is not None else text


def clean_table(table):
    rows = [[(" ".join(str(c).split()) if c else "") for c in row] for row in table]
    rows = [r for r in rows if any(r)]
    if not rows:
        return []
    keep = [i for i in range(max(len(r) for r in rows)) if any(i < len(r) and r[i] for r in rows)]
    return [[r[i] if i < len(r) else "" for i in keep] for r in rows]


def extract(pdf_path, page_spec):
    tables, text_pages, total = [], 0, 0
    with pdfplumber.open(pdf_path) as pdf:
        total = len(pdf.pages)
        for idx in parse_pages(page_spec, total):
            page = pdf.pages[idx]
            if (page.extract_text() or "").strip():
                text_pages += 1
            found = page.extract_tables() or page.extract_tables(TEXT_STRATEGY)
            for t in found:
                cleaned = clean_table(t)
                if len(cleaned) >= 2:
                    tables.append({"page": idx + 1, "rows": cleaned})
    return tables, text_pages, total


def merge_tables(tables):
    """Join consecutive tables with the same width, dropping repeated headers."""
    merged = []
    for t in tables:
        prev = merged[-1] if merged else None
        if prev and len(prev["rows"][0]) == len(t["rows"][0]):
            rows = t["rows"][1:] if t["rows"][0] == prev["rows"][0] else t["rows"]
            prev["rows"].extend(rows)
            prev["pages"].append(t["page"])
        else:
            merged.append({"pages": [t["page"]], "rows": [list(r) for r in t["rows"]]})
    return merged


def write_sheet(ws, rows, decimal):
    header_fill = PatternFill("solid", fgColor="DDE7F3")
    widths = {}
    for r_idx, row in enumerate(rows, start=1):
        for c_idx, raw in enumerate(row, start=1):
            value = raw if r_idx == 1 else convert_cell(raw, decimal)
            cell = ws.cell(row=r_idx, column=c_idx, value=value)
            if r_idx == 1:
                cell.font = Font(bold=True)
                cell.fill = header_fill
                cell.alignment = Alignment(wrap_text=True, vertical="top")
            elif isinstance(value, float):
                if value.is_integer() and not re.search(r"[.,]\d{2}\D*$", str(raw)):
                    cell.value = int(value)  # years, IDs, counts: no decimals
                else:
                    cell.number_format = "#,##0.00"
            elif hasattr(value, "year"):
                cell.number_format = "yyyy-mm-dd"
            widths[c_idx] = max(widths.get(c_idx, 0), min(len(str(raw)), 60))
    for c_idx, width in widths.items():
        ws.column_dimensions[get_column_letter(c_idx)].width = max(width + 2, 8)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("inputs", nargs="+", help="PDF file(s)")
    ap.add_argument("-o", "--output", help="Output .xlsx (default: <first input>.xlsx)")
    ap.add_argument("--pages", help="Pages to read, e.g. 1-3,5 (default: all)")
    ap.add_argument("--decimal", choices=["auto", "comma", "dot"], default="auto",
                    help="Decimal separator in the PDF (default: auto-detect per cell)")
    ap.add_argument("--no-merge", action="store_true", help="Keep every table on its own sheet")
    ap.add_argument("--csv", action="store_true", help="Also write one CSV per sheet")
    args = ap.parse_args()

    output = Path(args.output or Path(args.inputs[0]).with_suffix(".xlsx"))
    wb = Workbook()
    wb.remove(wb.active)
    summary = {"output": str(output), "sheets": [], "warnings": []}

    for pdf_path in args.inputs:
        if not Path(pdf_path).is_file():
            summary["warnings"].append(f"{pdf_path}: file not found")
            continue
        tables, text_pages, total = extract(pdf_path, args.pages)
        if text_pages == 0:
            summary["warnings"].append(
                f"{pdf_path}: no text layer found - this looks like a scanned PDF and needs OCR first")
            continue
        if not tables:
            summary["warnings"].append(f"{pdf_path}: text found but no tables detected")
            continue
        groups = ([{"pages": [t["page"]], "rows": t["rows"]} for t in tables]
                  if args.no_merge else merge_tables(tables))
        stem = Path(pdf_path).stem[:20]
        for n, group in enumerate(groups, start=1):
            title = f"{stem}_{n}" if len(groups) > 1 or len(args.inputs) > 1 else "Data"
            title = re.sub(r"[\[\]:*?/\\]", "_", title)[:31]
            write_sheet(wb.create_sheet(title), group["rows"], args.decimal)
            summary["sheets"].append({
                "sheet": title, "source": pdf_path, "pages": sorted(set(group["pages"])),
                "rows": len(group["rows"]) - 1, "columns": len(group["rows"][0]),
                "header": group["rows"][0],
            })
            if args.csv:
                csv_path = output.with_name(f"{output.stem}_{title}.csv")
                with open(csv_path, "w", newline="", encoding="utf-8-sig") as fh:
                    csv.writer(fh).writerows(group["rows"])

    if not summary["sheets"]:
        print(json.dumps(summary, indent=2, ensure_ascii=False))
        sys.exit(1)
    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)
    print(json.dumps(summary, indent=2, ensure_ascii=False, default=str))


if __name__ == "__main__":
    main()
